from __future__ import annotations

import json
import os
import tempfile
from datetime import datetime
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import unquote, urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
DB_PATH = ROOT / "base-datos-planificador.xlsx"
STATE_SHEET = "Estado JSON"
HOST = "127.0.0.1"
PORT = 8765


def text(value):
    if value is None:
        return ""
    return str(value)


def autosize(ws):
    for column_cells in ws.columns:
        max_len = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            max_len = max(max_len, len(text(cell.value)))
        ws.column_dimensions[column_letter].width = min(max(max_len + 2, 10), 46)


def style_header(ws):
    fill = PatternFill("solid", fgColor="DDEBE4")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_sheet(ws, headers, rows):
    ws.append(headers)
    for row in rows:
        ws.append(row)
    style_header(ws)
    autosize(ws)


def save_workbook(state):
    guests = state.get("guests", [])
    tables = state.get("tables", [])
    room = state.get("room", {})
    elements = room.get("elements", [])
    saved_at = datetime.now().isoformat(timespec="seconds")

    guest_by_id = {guest.get("id"): guest for guest in guests}

    wb = Workbook()
    ws = wb.active
    ws.title = "Asignaciones"

    assignment_rows = []
    seated_ids = set()
    for table in tables:
        for index, guest_id in enumerate(table.get("seats", []), start=1):
            guest = guest_by_id.get(guest_id)
            if guest:
                seated_ids.add(guest_id)
            assignment_rows.append([
                saved_at,
                table.get("name", ""),
                index,
                guest.get("name", "") if guest else "",
                guest.get("allergy", "") if guest else "",
                table.get("shape", ""),
                table.get("capacity", ""),
                table.get("x", ""),
                table.get("y", ""),
                table.get("rotation", 0),
            ])

    for guest in guests:
        if guest.get("id") not in seated_ids:
            assignment_rows.append([
                saved_at,
                "Sin asignar",
                "",
                guest.get("name", ""),
                guest.get("allergy", ""),
                "",
                "",
                "",
                "",
                "",
            ])

    write_sheet(
        ws,
        ["Guardado", "Mesa", "Silla", "Invitado", "Alergia", "Forma", "Capacidad", "Mesa X %", "Mesa Y %", "Rotacion"],
        assignment_rows,
    )

    ws_guests = wb.create_sheet("Invitados")
    write_sheet(
        ws_guests,
        ["ID", "Nombre", "Alergia", "Sentado"],
        [[guest.get("id", ""), guest.get("name", ""), guest.get("allergy", ""), "Si" if guest.get("id") in seated_ids else "No"] for guest in guests],
    )

    ws_tables = wb.create_sheet("Mesas")
    write_sheet(
        ws_tables,
        ["ID", "Nombre", "Forma", "Capacidad", "Ocupadas", "X %", "Y %", "Rotacion"],
        [
            [
                table.get("id", ""),
                table.get("name", ""),
                table.get("shape", ""),
                table.get("capacity", ""),
                len([seat for seat in table.get("seats", []) if seat]),
                table.get("x", ""),
                table.get("y", ""),
                table.get("rotation", 0),
            ]
            for table in tables
        ],
    )

    ws_room = wb.create_sheet("Sala")
    door = room.get("door", {})
    write_sheet(
        ws_room,
        ["Guardado", "Forma", "Ancho m", "Alto m", "Zoom", "Puerta lado", "Puerta posicion"],
        [[saved_at, room.get("shape", ""), room.get("width", ""), room.get("height", ""), room.get("zoom", ""), door.get("side", ""), door.get("pos", "")]],
    )

    ws_elements = wb.create_sheet("Elementos")
    write_sheet(
        ws_elements,
        ["ID", "Elemento", "X %", "Y %", "Ancho m", "Alto m", "Rotacion"],
        [[item.get("id", ""), item.get("label", ""), item.get("x", ""), item.get("y", ""), item.get("w", ""), item.get("h", ""), item.get("rotation", 0)] for item in elements],
    )

    ws_state = wb.create_sheet(STATE_SHEET)
    state_json = json.dumps(state, ensure_ascii=False)
    ws_state.append(["Parte", "JSON"])
    chunk_size = 30000
    for index in range(0, len(state_json), chunk_size):
        ws_state.append([index // chunk_size + 1, state_json[index:index + chunk_size]])
    ws_state.sheet_state = "hidden"

    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx", dir=ROOT) as tmp:
        temp_path = Path(tmp.name)
    try:
        wb.save(temp_path)
        os.replace(temp_path, DB_PATH)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def cell_value(ws, row, column, default=""):
    value = ws.cell(row=row, column=column).value
    return default if value is None else value


def load_state_from_json_sheet(wb):
    if STATE_SHEET not in wb.sheetnames:
        return None
    ws = wb[STATE_SHEET]
    chunks = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[1]:
            chunks.append(str(row[1]))
    if not chunks:
        return None
    return json.loads("".join(chunks))


def load_state_from_visible_sheets(wb):
    guests = []
    guest_by_name = {}
    guest_by_id = {}
    if "Invitados" in wb.sheetnames:
        ws = wb["Invitados"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[1]:
                continue
            guest = {"id": text(row[0]), "name": text(row[1]), "allergy": text(row[2]) if len(row) > 2 else ""}
            guests.append(guest)
            guest_by_id[guest["id"]] = guest
            guest_by_name[guest["name"]] = guest

    tables = []
    table_by_name = {}
    if "Mesas" in wb.sheetnames:
        ws = wb["Mesas"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[1]:
                continue
            capacity = int(row[3] or 0)
            table = {
                "id": text(row[0]),
                "name": text(row[1]),
                "shape": text(row[2]) or "round",
                "capacity": capacity,
                "x": row[5] if len(row) > 5 and row[5] is not None else 50,
                "y": row[6] if len(row) > 6 and row[6] is not None else 50,
                "rotation": row[7] if len(row) > 7 and row[7] is not None else 0,
                "seats": [None] * capacity,
            }
            tables.append(table)
            table_by_name[table["name"]] = table

    if "Asignaciones" in wb.sheetnames:
        ws = wb["Asignaciones"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[1] == "Sin asignar" or not row[2] or not row[3]:
                continue
            table = table_by_name.get(text(row[1]))
            guest = guest_by_name.get(text(row[3]))
            if not table or not guest:
                continue
            seat_index = int(row[2]) - 1
            if 0 <= seat_index < len(table["seats"]):
                table["seats"][seat_index] = guest["id"]

    room = {
        "shape": "rect",
        "width": 18,
        "height": 12,
        "zoom": 1,
        "door": {"side": "bottom", "pos": .5},
        "elements": [],
    }
    if "Sala" in wb.sheetnames and wb["Sala"].max_row >= 2:
        ws = wb["Sala"]
        room.update({
            "shape": cell_value(ws, 2, 2, "rect"),
            "width": cell_value(ws, 2, 3, 18),
            "height": cell_value(ws, 2, 4, 12),
            "zoom": cell_value(ws, 2, 5, 1),
            "door": {"side": cell_value(ws, 2, 6, "bottom"), "pos": cell_value(ws, 2, 7, .5)},
        })

    if "Elementos" in wb.sheetnames:
        ws = wb["Elementos"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[1]:
                continue
            room["elements"].append({
                "id": text(row[0]),
                "label": text(row[1]),
                "x": row[2] if len(row) > 2 and row[2] is not None else 50,
                "y": row[3] if len(row) > 3 and row[3] is not None else 50,
                "w": row[4] if len(row) > 4 and row[4] is not None else 2.4,
                "h": row[5] if len(row) > 5 and row[5] is not None else 1.1,
                "rotation": row[6] if len(row) > 6 and row[6] is not None else 0,
            })

    return {"guests": guests, "tables": tables, "room": room, "mapCollapsed": False, "mapFullscreen": False, "mapPanelWidth": 640}


def load_workbook_state():
    if not DB_PATH.exists():
        return None
    wb = load_workbook(DB_PATH, data_only=True)
    try:
        state = load_state_from_json_sheet(wb)
        if state:
            return state
        return load_state_from_visible_sheets(wb)
    finally:
        wb.close()


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == "/api/load":
            try:
                state = load_workbook_state()
                body = json.dumps({"ok": True, "exists": state is not None, "state": state}, ensure_ascii=False).encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
            except Exception as exc:
                body = json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False).encode("utf-8")
                self.send_response(500)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
            return
        super().do_GET()

    def do_POST(self):
        parsed = urlparse(self.path)
        if parsed.path != "/api/save":
            self.send_error(404)
            return

        try:
            length = int(self.headers.get("Content-Length", "0"))
            payload = self.rfile.read(length).decode("utf-8")
            state = json.loads(payload)
            save_workbook(state)
            body = json.dumps({"ok": True, "path": str(DB_PATH), "savedAt": datetime.now().isoformat(timespec="seconds")}).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
        except Exception as exc:
            body = json.dumps({"ok": False, "error": str(exc)}).encode("utf-8")
            self.send_response(500)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

    def translate_path(self, path):
        parsed = urlparse(path)
        clean = unquote(parsed.path).lstrip("/")
        if not clean:
            clean = "planificador-mesas-boda.html"
        return str(ROOT / clean)


if __name__ == "__main__":
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    print(f"Planificador listo en http://{HOST}:{PORT}/planificador-mesas-boda.html")
    print(f"Base de datos Excel: {DB_PATH}")
    server.serve_forever()
